from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SECTOR_THEMES = {
    "marketing": {"accent": "2FA7A0", "navy": "16324F", "highlight": "2E8B70", "background": "F7F9FB"},
    "sales": {"accent": "2FA7A0", "navy": "16324F", "highlight": "2E8B70", "background": "F7F9FB"},
    "supply": {"accent": "2FA7A0", "navy": "16324F", "highlight": "2E8B70", "background": "F7F9FB"},
    "default": {"accent": "2FA7A0", "navy": "16324F", "highlight": "2E8B70", "background": "F7F9FB"},
}

MUTED = "64748B"


def _theme_for_product(product_id: str) -> dict:
    if product_id in ("product-10", "product-11", "product-12"):
        return SECTOR_THEMES["supply"]
    if product_id in ("product-02", "product-13", "product-14"):
        return SECTOR_THEMES["sales"]
    if product_id in ("product-03", "product-04", "product-05", "product-06", "product-07", "product-08", "product-09"):
        return SECTOR_THEMES["marketing"]
    return SECTOR_THEMES["default"]


def _color(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def _add_text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _color(color)
    return box


def _set_background(slide, hex_color: str):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _color(hex_color)


def build_decision_brief_pptx(product_id: str, status: str, output: dict) -> bytes:
    palette = _theme_for_product(product_id)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def footer(slide):
        _add_text(slide, 0.5, 7.05, 3.2, 0.3, "Made by PharmaLens AI", 8, MUTED, bold=True)
        _add_text(slide, 8.6, 7.05, 3.9, 0.3, "Evidence first / Decisions forward", 8, MUTED, align=PP_ALIGN.RIGHT)

    # Slide 1 — cover
    cover = prs.slides.add_slide(blank)
    _set_background(cover, palette["background"])
    _add_text(cover, 0.65, 0.55, 4, 0.3, "PHARMALENS AI", 10, palette["accent"], bold=True)
    _add_text(cover, 0.65, 1.55, 8.5, 0.8, "Decision brief", 32, palette["navy"], bold=True)
    _add_text(cover, 0.65, 2.35, 8.5, 0.6, product_id, 25, palette["accent"], bold=True)
    _add_text(cover, 0.65, 3.25, 7.7, 1.2, output.get("summary") or "Saved pharmaceutical decision output.", 16, palette["navy"])
    _add_text(cover, 0.65, 5.6, 3.5, 0.3, f"Status: {status}", 11, MUTED, bold=True)
    footer(cover)

    # Slide 2 — decision signals
    metrics_slide = prs.slides.add_slide(blank)
    _set_background(metrics_slide, palette["background"])
    _add_text(metrics_slide, 0.65, 0.55, 4, 0.3, "01 / Decision signals", 10, palette["accent"], bold=True)
    _add_text(metrics_slide, 0.65, 1.1, 8, 0.6, "What the data says", 26, palette["navy"], bold=True)
    entries = list((output.get("metrics") or {}).items())[:6]
    for index, (key, value) in enumerate(entries):
        col, row = index % 3, index // 3
        x, y = 0.65 + col * 3.45, 2.0 + row * 1.55
        _add_text(metrics_slide, x, y, 2.7, 0.25, str(key).replace("_", " "), 10, MUTED, bold=True)
        _add_text(metrics_slide, x, y + 0.3, 2.7, 0.5, str(value if value is not None else "\u2014"), 20, palette["navy"], bold=True)
    footer(metrics_slide)

    # Slide 3 — evidence & confidence
    evidence_slide = prs.slides.add_slide(blank)
    _set_background(evidence_slide, "111111")
    _add_text(evidence_slide, 0.65, 0.55, 5, 0.3, "02 / Evidence & confidence", 10, palette["highlight"], bold=True)
    _add_text(evidence_slide, 0.65, 1.1, 7, 0.6, "Show your work.", 26, "FFFFFF", bold=True)
    confidence = output.get("confidence") or {}
    score = round((confidence.get("score") or 0) * 100)
    _add_text(evidence_slide, 7.8, 1.05, 2.2, 0.7, f"{score}%", 32, palette["highlight"], bold=True, align=PP_ALIGN.RIGHT)
    _add_text(evidence_slide, 7.8, 1.75, 2.2, 0.3, confidence.get("level") or "unknown", 10, "FFFFFF", bold=True, align=PP_ALIGN.RIGHT)
    for index, item in enumerate((output.get("evidence") or [])[:6]):
        y = 2.15 + index * 0.55
        _add_text(evidence_slide, 0.65, y, 8.8, 0.3, f"{item.get('field')}: {item.get('value') if item.get('value') is not None else '-'}", 12, "FFFFFF")
        _add_text(evidence_slide, 9.6, y, 3.0, 0.3, str(item.get("source", "")), 9, palette["highlight"], align=PP_ALIGN.RIGHT)
    _add_text(evidence_slide, 0.65, 5.95, 8.8, 0.5, confidence.get("rationale") or "Review the evidence before making a business decision.", 11, "C9C9C9")
    footer(evidence_slide)


    # Visualization slides — use the same saved ChartSpec contract as the web UI.
    for chart in (output.get("visualizations") or [])[:6]:
        slide = prs.slides.add_slide(blank)
        _set_background(slide, palette["background"])
        title = (chart.get("meta") or {}).get("title") or chart.get("chartId") or "Visualization"
        description = (chart.get("meta") or {}).get("description") or ""
        role = (chart.get("presentationPriority") or {}).get("role", "diagnostic")
        quality = chart.get("presentationQuality") or {}
        label = "PRIMARY VISUAL" if role == "primary" else "DIAGNOSTIC VISUAL"
        _add_text(slide, 0.65, 0.55, 4, 0.3, label, 10, palette["accent"], bold=True)
        if quality.get("warnings"):
            _add_text(slide, 9.4, 0.55, 2.6, 0.3, "REVIEW VISUAL DENSITY", 8, "C94C4C", bold=True, align=PP_ALIGN.RIGHT)
        _add_text(slide, 0.65, 1.05, 10.8, 0.55, str(title), 24, palette["navy"], bold=True)
        if description:
            _add_text(slide, 0.65, 1.62, 10.8, 0.4, str(description), 10, MUTED)
        data = chart.get("data") or []
        series = chart.get("series") or []
        x_key = chart.get("xKey")
        primary = series[0] if series else {}
        data_key = primary.get("dataKey")
        values = []
        for row in data[:8]:
            try:
                value = float(row.get(data_key, 0)) if data_key else 0
            except (TypeError, ValueError):
                value = 0
            values.append((str(row.get(x_key, "")) if x_key else "", value))
        max_value = max((abs(v) for _, v in values), default=1) or 1
        for index, (label, value) in enumerate(values):
            y = 2.25 + index * 0.52
            _add_text(slide, 0.65, y, 2.4, 0.25, label[:32], 9, MUTED)
            bar_width = 6.4 * abs(value) / max_value
            shape = slide.shapes.add_shape(1, Inches(3.05), Inches(y), Inches(max(0.08, bar_width)), Inches(0.22))
            shape.fill.solid()
            shape.fill.fore_color.rgb = _color(palette["accent"])
            shape.line.fill.background()
            _add_text(slide, 9.7, y, 2.2, 0.25, f"{value:,.1f}", 9, palette["navy"], bold=True, align=PP_ALIGN.RIGHT)
        footer(slide)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
