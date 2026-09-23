"""Professional local PowerPoint renderer with sector themes and PharmaLens branding."""
from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

try:
    from pharmalens.copilot.limits import UsageLimiter
except ImportError:
    UsageLimiter = None  # legacy optional dependency
from .slide_spec import PresentationSpec

BRAND = "Made by PharmaLens AI"


@dataclass(frozen=True)
class Theme:
    name: str
    navy: str
    accent: str
    highlight: str
    background: str
    light: str


THEMES = {
    "marketing": Theme("Marketing", "173F5F", "20639B", "F6D55C", "F5F8FC", "E8F1FA"),
    "sales": Theme("Sales", "12372A", "2E8B57", "F4A261", "F1F8F4", "DFF2E5"),
    "supply": Theme("Supply", "263238", "00838F", "FFB703", "F3F7F8", "DDF3F4"),
    "default": Theme("PharmaLens", "102A43", "1D70B8", "54C2FF", "F5F8FC", "E7F1FA"),
}


def get_theme(theme: str = "default") -> Theme:
    return THEMES.get(theme.lower(), THEMES["default"])


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def render_pptx(spec: PresentationSpec, output_path: str | Path, *, limiter: Any | None = None, theme: str = "default") -> Path:
    if limiter is None and UsageLimiter is not None:
        limiter = UsageLimiter.from_env()
    max_slides = int(os.getenv("PPT_MAX_SLIDES_FREE", "6"))
    spec.validate(max_slides=max_slides)
    if limiter is not None and not limiter.allow("ppt"):
        raise RuntimeError("Daily PPT limit reached")
    palette = get_theme(theme)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for index, item in enumerate(spec.slides):
        slide = prs.slides.add_slide(blank)
        _background(slide, palette)
        if index == 0:
            _title_slide(slide, spec.title, spec.subtitle, palette)
        else:
            _header(slide, item.title, palette)
            if item.slide_type == "kpi":
                _render_kpis(slide, item.metrics, palette)
            elif item.slide_type == "bar_chart":
                _render_chart(slide, item.metrics, palette)
            else:
                _render_bullets(slide, item.bullets, palette)
        _footer(slide, palette, index + 1, len(spec.slides))
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    if limiter is not None:
        limiter.record("ppt")
    return path


def _background(slide: Any, palette: Theme) -> None:
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = _rgb(palette.background)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(.16))
    band.fill.solid(); band.fill.fore_color.rgb = _rgb(palette.accent); band.line.fill.background()


def _title_slide(slide: Any, title: str, subtitle: str, palette: Theme) -> None:
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.7), Inches(1.1), Inches(.18), Inches(4.3))
    accent.fill.solid(); accent.fill.fore_color.rgb = _rgb(palette.highlight); accent.line.fill.background()
    box = slide.shapes.add_textbox(Inches(1.2), Inches(1.25), Inches(11), Inches(2.2))
    frame = box.text_frame; frame.clear()
    p = frame.paragraphs[0]; p.text = title; p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = _rgb(palette.navy)
    p2 = frame.add_paragraph(); p2.text = subtitle; p2.font.size = Pt(17); p2.font.color.rgb = _rgb(palette.accent); p2.space_before = Pt(18)
    mark = slide.shapes.add_textbox(Inches(1.2), Inches(4.65), Inches(6), Inches(.5))
    mark.text_frame.text = BRAND; mark.text_frame.paragraphs[0].font.size = Pt(15); mark.text_frame.paragraphs[0].font.bold = True; mark.text_frame.paragraphs[0].font.color.rgb = _rgb(palette.accent)


def _header(slide: Any, title: str, palette: Theme) -> None:
    box = slide.shapes.add_textbox(Inches(.7), Inches(.55), Inches(11.8), Inches(.65))
    p = box.text_frame.paragraphs[0]; p.text = title; p.font.size = Pt(25); p.font.bold = True; p.font.color.rgb = _rgb(palette.navy)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.7), Inches(1.25), Inches(1.05), Inches(.07))
    rule.fill.solid(); rule.fill.fore_color.rgb = _rgb(palette.highlight); rule.line.fill.background()


def _footer(slide: Any, palette: Theme, number: int, total: int) -> None:
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.7), Inches(6.82), Inches(11.95), Inches(.012))
    line.fill.solid(); line.fill.fore_color.rgb = _rgb(palette.light); line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(.7), Inches(6.9), Inches(9), Inches(.25))
    p = box.text_frame.paragraphs[0]; p.text = BRAND; p.font.size = Pt(8); p.font.bold = True; p.font.color.rgb = _rgb(palette.accent)
    page = slide.shapes.add_textbox(Inches(11.8), Inches(6.9), Inches(.8), Inches(.25))
    p = page.text_frame.paragraphs[0]; p.text = f"{number}/{total}"; p.font.size = Pt(8); p.alignment = PP_ALIGN.RIGHT; p.font.color.rgb = _rgb(palette.navy)


def _render_kpis(slide: Any, metrics: list[dict[str, Any]], palette: Theme) -> None:
    width = 11.5 / max(1, min(4, len(metrics)))
    for i, metric in enumerate(metrics[:4]):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.7 + i * width), Inches(2.0), Inches(width - .25), Inches(2.0))
        card.fill.solid(); card.fill.fore_color.rgb = _rgb(palette.light); card.line.color.rgb = _rgb(palette.accent)
        box = slide.shapes.add_textbox(Inches(.95 + i * width), Inches(2.35), Inches(width - .7), Inches(1.3))
        frame = box.text_frame; frame.clear()
        p = frame.paragraphs[0]; p.text = str(metric.get("label", "Metric")); p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = _rgb(palette.navy)
        p2 = frame.add_paragraph(); p2.text = str(metric.get("value", "")); p2.font.size = Pt(24); p2.font.bold = True; p2.font.color.rgb = _rgb(palette.accent); p2.space_before = Pt(10)


def _render_chart(slide: Any, metrics: list[dict[str, Any]], palette: Theme) -> None:
    data = CategoryChartData(); data.categories = [str(x.get("label", "")) for x in metrics]; data.add_series("Value", [float(x.get("value", 0)) for x in metrics])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(1.65), Inches(11), Inches(4.8), data).chart
    chart.has_legend = False; chart.has_title = False
    chart.value_axis.tick_labels.font.size = Pt(9); chart.category_axis.tick_labels.font.size = Pt(10)
    chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb = _rgb(palette.accent)


def _render_bullets(slide: Any, bullets: list[str], palette: Theme) -> None:
    box = slide.shapes.add_textbox(Inches(1), Inches(1.7), Inches(11), Inches(4.6))
    frame = box.text_frame; frame.clear()
    for i, text in enumerate(bullets):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph(); p.text = text; p.font.size = Pt(19); p.font.color.rgb = _rgb(palette.navy); p.level = 0; p.space_after = Pt(13)
